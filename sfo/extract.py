"""正文抽取：本地、只读、尽力而为。

原则：
- 任何格式失败都降级为空文本 + warning，绝不中断整批归档；
- 只读前若干页/行，并限制文件大小，避免被超大文件拖死；
- 不修改源文件，不调用外部服务。
"""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass

from .textutil import collapse_ws, decode_bytes

TEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".csv", ".tsv", ".json", ".xml", ".ini"}
OFFICE_EXTS = {".docx", ".xlsx", ".xlsm", ".pptx"}
PDF_EXTS = {".pdf"}
SUPPORTED_EXTS = TEXT_EXTS | OFFICE_EXTS | PDF_EXTS


@dataclass
class Extracted:
    text: str = ""
    source: str = "none"
    warning: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.text)


def extract_text(path: str, max_chars: int = 1500, max_file_mb: int = 50) -> Extracted:
    ext = os.path.splitext(path)[1].lower()
    try:
        size = os.path.getsize(path)
    except OSError as exc:
        return Extracted(warning=f"无法读取文件信息：{exc}")
    if size == 0:
        return Extracted(warning="空文件")
    if size > max_file_mb * 1024 * 1024:
        return Extracted(warning=f"文件超过 {max_file_mb}MB，仅使用文件名")

    try:
        if ext in TEXT_EXTS:
            result = _extract_text_file(path, ext)
        elif ext == ".docx":
            result = _extract_docx(path)
        elif ext in (".xlsx", ".xlsm"):
            result = _extract_xlsx(path)
        elif ext == ".pptx":
            result = _extract_pptx(path)
        elif ext in PDF_EXTS:
            result = _extract_pdf(path)
        else:
            return Extracted(source="none", warning=f"暂不支持 {ext or '无扩展名'} 的正文抽取")
    except ImportError as exc:
        return Extracted(source="none", warning=f"缺少可选依赖：{exc}")
    except Exception as exc:  # 抽取永远不能让归档流程失败
        return Extracted(source="none", warning=f"抽取失败：{type(exc).__name__}: {exc}")

    text = collapse_ws(result.text)[:max_chars]
    result.text = text
    return result


def _extract_text_file(path: str, ext: str) -> Extracted:
    with open(path, "rb") as f:
        raw = f.read(256 * 1024)
    text = decode_bytes(raw)
    if ext in (".csv", ".tsv"):
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
            rows = list(csv.reader(text.splitlines()[:30], dialect))
            text = " ".join(" ".join(cell for cell in row[:10]) for row in rows[:20])
        except (csv.Error, TypeError):
            pass
    return Extracted(text=text, source=ext.lstrip("."))


def _extract_docx(path: str) -> Extracted:
    from docx import Document  # python-docx

    document = Document(path)
    parts: list[str] = []
    for paragraph in document.paragraphs[:60]:
        value = paragraph.text.strip()
        if value:
            parts.append(value)
    for table in document.tables[:3]:
        for row in table.rows[:5]:
            cells = [cell.text.strip() for cell in row.cells[:8] if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return Extracted(text="\n".join(parts), source="docx")


def _extract_xlsx(path: str) -> Extracted:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets[:3]:
            parts.append(sheet.title)
            for index, row in enumerate(sheet.iter_rows(max_row=30, max_col=10, values_only=True)):
                if index >= 30:
                    break
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    parts.append(" | ".join(cells[:10]))
    finally:
        workbook.close()
    return Extracted(text="\n".join(parts), source="xlsx")


def _extract_pptx(path: str) -> Extracted:
    from pptx import Presentation

    presentation = Presentation(path)
    parts: list[str] = []
    for slide in presentation.slides[:15]:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            value = shape.text_frame.text.strip()
            if value:
                parts.append(value)
    return Extracted(text="\n".join(parts), source="pptx")


def _extract_pdf(path: str) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # 兼容旧环境
        except ImportError as exc:
            raise ImportError("pypdf（pip install pypdf）用于读取 PDF 正文") from exc
    reader = PdfReader(path)
    parts: list[str] = []
    for page in reader.pages[:5]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(parts)
    if not text.strip():
        return Extracted(
            text="", source="pdf", warning="PDF 没有可抽取的文本层（可能是扫描件，需 OCR）"
        )
    return Extracted(text=text, source="pdf")
