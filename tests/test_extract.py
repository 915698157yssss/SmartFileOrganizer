import os
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from sfo.extract import extract_text
from tutil import cleanup, make_temp_dir


class ExtractTest(unittest.TestCase):
    def setUp(self):
        self.tmp_path = make_temp_dir("extract")
        self.tmp = str(self.tmp_path)

    def tearDown(self):
        cleanup(self.tmp_path)

    def test_txt(self):
        path = os.path.join(self.tmp, "a.txt")
        Path(path).write_text("设备验收报告\n请于2026年12月31日前反馈", encoding="utf-8")
        result = extract_text(path)
        self.assertTrue(result.ok)
        self.assertIn("验收", result.text)
        self.assertEqual(result.source, "txt")

    def test_gbk_txt(self):
        path = os.path.join(self.tmp, "b.txt")
        Path(path).write_bytes("会议纪要 参会人员".encode("gbk"))
        self.assertIn("会议纪要", extract_text(path).text)

    def test_docx(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx 未安装")
        path = os.path.join(self.tmp, "c.docx")
        document = Document()
        document.add_paragraph("一期竣工验收报告")
        document.add_paragraph("请于2026年11月30日前完成整改")
        document.save(path)
        result = extract_text(path)
        self.assertTrue(result.ok)
        self.assertIn("竣工", result.text)

    def test_xlsx(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl 未安装")
        path = os.path.join(self.tmp, "d.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "结算单"
        sheet.append(["项目", "金额"])
        sheet.append(["一期工程", "286500元"])
        workbook.save(path)
        result = extract_text(path)
        self.assertTrue(result.ok)
        self.assertIn("结算单", result.text)

    def test_unsupported_degrades(self):
        path = os.path.join(self.tmp, "e.bin")
        Path(path).write_bytes(b"\x00\x01\x02")
        result = extract_text(path)
        self.assertFalse(result.ok)
        self.assertTrue(result.warning)


if __name__ == "__main__":
    unittest.main()
